from copy import deepcopy
import json

import pandas as pd
import pytest

from autobot import tender_corrections as edit, estimate_publication as publication, estimate_parse_worker as worker
from autobot import main, estimate_publication_recovery as recovery
from autobot.market_contract import match_market_rows
from test_estimate_parse_pipeline import setup, TID, saved_reports

ACTOR = {'id': 7, 'name': 'Автор проверки'}


@pytest.fixture
def context(tmp_path, monkeypatch):
    paths, source, tender = setup(tmp_path)
    monkeypatch.setattr(publication, 'run_parser', lambda excel, pdf, item: worker.parse_files(excel, pdf, item, worker.ParseLimits()))
    publication.parse_and_publish(tender, [source], [], [source], paths)
    return paths, source, tender


def data(context, **updates):
    paths, _, _ = context
    value = edit.snapshot(paths['reports'], TID)
    body = {'position_id': value['rows'][0]['position_id'], 'changes': {'unit_price': '125.05', 'total': '25.01'},
            'expected_version': value['version'], 'operation_id': 'b'*32, 'reason': 'Сверено с исходным документом'}
    body.update(updates)
    return body


def save(context, **updates):
    paths, _, tender = context
    return edit.apply(tender, paths, actor=ACTOR, **data(context, **updates))


def test_canonical_excel_history_original_sources_and_market(context):
    paths, source, _ = context
    source_bytes = source.read_bytes()
    before = edit.snapshot(paths['reports'], TID)
    first, duplicate = save(context)
    current = edit.snapshot(paths['reports'], TID)
    assert not duplicate and current['revision'] == 1 and current['version'] == first['version']
    assert current['rows'][0]['total'] == '25.01' and current['rows'][0]['unit_price'] == '125.05'
    assert int(current['frame'].iloc[0]['total_kopecks']) == 2501
    assert int(current['frame'].iloc[0]['unit_price_kopecks']) == 12505
    assert current['rows'][1] == before['rows'][1] and current['original_rows'] == before['rows']
    assert source.read_bytes() == source_bytes
    assert current['manifest']['parse_sources'] == before['manifest']['parse_sources']
    assert current['manifest']['parse_resources'] == before['manifest']['parse_resources']
    assert current['manifest']['official_total_rub'] == before['manifest']['official_total_rub']
    assert edit.history(current)[0]['changes']['total']['after'] == '25.01'
    matches = match_market_rows(current['frame'], before['frame'])
    assert matches[0] is None and matches[1] is not None
    assert '25,01' in (paths['reports']/recovery.output_names(TID)[0]).read_text(encoding='utf-8')


def test_repeat_after_later_edit_and_different_actor_conflict(context):
    paths, _, tender = context
    body = data(context)
    first, _ = edit.apply(tender, paths, actor=ACTOR, **body)
    save(context, changes={'qty': '0.2'}, operation_id='c'*32)
    result, duplicate = edit.apply(tender, paths, actor=dict(ACTOR, name='Новое имя'), **body)
    assert duplicate and result == first and edit.snapshot(paths['reports'], TID)['revision'] == 2
    for args in ({'reason': 'другая причина'}, {'changes': {'total': '1'}}):
        with pytest.raises(edit.CorrectionError) as error:
            edit.apply(tender, paths, actor=ACTOR, **dict(body, **args))
        assert error.value.status == 409
    with pytest.raises(edit.CorrectionError):
        edit.apply(tender, paths, actor={'id': 8, 'name': 'Другой'}, **body)
    current = edit.snapshot(paths['reports'], TID)
    assert edit.receipt(current, body['operation_id'], ACTOR)['revision'] == 1
    assert edit.receipt(current, body['operation_id'], {'id': 8}) is None


def test_unknown_and_zero_remain_distinct_without_automatic_product(context):
    paths, _, _ = context
    save(context, changes={'qty': '', 'unit_price': '0', 'total': ''})
    current = edit.snapshot(paths['reports'], TID)
    row = current['rows'][0]
    assert row['qty'] is None and row['unit_price'] == '0' and row['total'] is None
    assert pd.isna(current['frame'].iloc[0]['total_kopecks'])
    save(context, changes={'total': '0'}, operation_id='c'*32)
    assert edit.snapshot(paths['reports'], TID)['rows'][0]['total'] == '0'


@pytest.mark.parametrize('changes', [{'total':'1.001'}, {'qty':'NaN'}, {'total':True}, {'qty':'0.0000001'}, {'name':''}, {'type':'work'}])
def test_invalid_values_do_not_change_report(context, changes):
    paths, _, _ = context
    before = saved_reports(paths)
    with pytest.raises(edit.CorrectionError):
        save(context, changes=changes)
    assert saved_reports(paths) == before


def test_source_change_and_stale_tab_preserve_current_values(context):
    paths, source, tender = context
    body = data(context)
    save(context)
    with pytest.raises(edit.CorrectionError) as error:
        edit.apply(tender, paths, actor=ACTOR, **dict(body, operation_id='c'*32))
    assert error.value.status == 409
    before = saved_reports(paths)
    source.write_bytes(b'changed')
    with pytest.raises(edit.CorrectionError):
        save(context, changes={'qty':'2'}, operation_id='d'*32)
    assert saved_reports(paths) == before


def test_corrupted_history_and_changed_excel_fail_closed(context):
    paths, _, _ = context
    save(context)
    manifest_path = paths['reports']/recovery.output_names(TID)[2]
    original = manifest_path.read_bytes()
    manifest = json.loads(original)
    manifest[edit.LEDGER_KEY]['events'][0]['reason'] = 'подмена истории'
    manifest_path.write_text(json.dumps(manifest,ensure_ascii=False),encoding='utf-8')
    with pytest.raises(edit.CorrectionError) as error:
        edit.snapshot(paths['reports'], TID)
    assert error.value.status == 503
    manifest_path.write_bytes(original)
    report = paths['reports']/recovery.output_names(TID)[1]
    frame = pd.read_excel(report)
    frame.loc[0,edit.COLUMNS['total']] = 123
    frame.to_excel(report,index=False)
    with pytest.raises(edit.CorrectionError):
        edit.snapshot(paths['reports'], TID)


def test_failed_second_replace_rolls_back_then_same_request_succeeds(context, monkeypatch):
    paths, _, tender = context
    body = data(context)
    before = saved_reports(paths)
    original_replace = recovery.os.replace
    failed = False
    def fail_once(source, destination):
        nonlocal failed
        if not failed and str(destination).endswith(recovery.output_names(TID)[1]):
            failed = True
            raise OSError('injected publication failure')
        return original_replace(source, destination)
    with monkeypatch.context() as patch:
        patch.setattr(recovery.os,'replace',fail_once)
        with pytest.raises(OSError):
            edit.apply(tender, paths, actor=ACTOR, **body)
    assert failed and saved_reports(paths) == before and edit.snapshot(paths['reports'], TID)['revision'] == 0
    result, duplicate = edit.apply(tender, paths, actor=ACTOR, **body)
    assert result['revision'] == 1 and not duplicate


def test_reparse_identical_base_preserves_correction_and_original_history(context):
    paths, source, tender = context
    save(context)
    before = edit.snapshot(paths['reports'], TID)
    publication.parse_and_publish(tender,[source],[],[source],paths)
    current = edit.snapshot(paths['reports'], TID)
    assert current['rows'] == before['rows'] and current['original_rows'] == before['original_rows']
    assert current['version'] == before['version'] and current['revision'] == 1


def test_reparse_different_base_keeps_canonical_and_history(context, monkeypatch):
    paths, source, tender = context
    save(context)
    before = saved_reports(paths)
    previous = publication.run_parser
    def different(*args):
        result = previous(*args)
        result['rows'][0]['work_name'] += ' изменено'
        return result
    monkeypatch.setattr(publication,'run_parser',different)
    with pytest.raises(worker.EstimateParseRejected,match='ручных исправлений'):
        publication.parse_and_publish(tender,[source],[],[source],paths)
    assert saved_reports(paths) == before
    assert edit.snapshot(paths['reports'], TID)['revision'] == 1


def test_excel_amounts_are_numeric_and_text_is_not_a_formula(context):
    from openpyxl import load_workbook
    paths, _, _ = context
    save(context,changes={'name':'=1+1','total':'25.01','unit_price':'125.05','qty':'0.2'})
    book=load_workbook(paths['reports']/recovery.output_names(TID)[1],data_only=False)
    try:
        sheet=book.active
        columns={cell.value:cell.column for cell in sheet[1]}
        assert sheet.cell(2,columns[edit.COLUMNS['name']]).data_type=='s'
        assert sheet.cell(2,columns[edit.COLUMNS['name']]).value=='=1+1'
        for name in ('qty','unit_price','total'):
            assert sheet.cell(2,columns[edit.COLUMNS[name]]).data_type=='n'
        assert sheet.cell(2,columns[edit.COLUMNS['total']]).value==25.01
    finally:book.close()


def test_history_removal_does_not_allow_reparse_to_erase_corrections(context):
    paths,source,tender=context
    save(context)
    manifest=paths['reports']/recovery.output_names(TID)[2]
    value=json.loads(manifest.read_text(encoding='utf-8'));value.pop(edit.LEDGER_KEY)
    manifest.write_text(json.dumps(value),encoding='utf-8')
    with pytest.raises(worker.EstimateParseRejected,match='отсутствует история'):
        publication.parse_and_publish(tender,[source],[],[source],paths)
    with pytest.raises(edit.CorrectionError,match='отсутствует история'):
        edit.snapshot(paths['reports'],TID)
